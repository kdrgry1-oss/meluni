"""
PAZARLAMA İZİNLERİ (E-posta & SMS) — kim izin verdi / kim reddetti, tek listede.

Kaynaklar (birleştirilir; alıcı+kanal başına EN SON kayıt geçerli):
  • db.iys_consents          — sipariş/üyelik/OTP akışlarından İYS izin kayıtları
                               (channels: MESAJ=SMS, EPOSTA=e-posta; status: ONAY / RET)
  • db.newsletter_subscribers — site altı bülten aboneliği (e-posta; consent=True, active)
  • db.users.accepts_marketing — üye profili pazarlama tercihi (Hesabım / admin / sipariş izni)
  • db.email_suppressions    — bounce/şikâyet kara listesi (e-posta izni fiilen geçersiz)

GET  /admin/consents?channel=email|sms&status=onay|ret&q=&page=&limit=
GET  /admin/consents/summary
GET  /admin/consents/export.xlsx
"""
import io
import re
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse

from .deps import db, require_admin, require_permission
from activity_audit import record_admin_audit

router = APIRouter(prefix="/admin/consents", tags=["admin-consents"])

_SRC_TR = {"HS_WEB": "Web sitesi", "HS_FIZIKSEL_ORTAM": "Fiziksel ortam", "HS_MESAJ": "SMS ile", "HS_EPOSTA": "E-posta ile",
           "footer": "Bülten (site altı)", "checkout": "Ödeme sayfası", "register": "Üyelik", "otp": "SMS doğrulama",
           "popup": "Açılır pencere", "account": "Hesabım", "IYS": "İYS (müşteri İYS'den reddetti)"}


async def _collect() -> dict:
    """{(kanal, alıcı): kayıt} — en güncel kayıt kazanır."""
    latest: dict = {}

    def _put(channel, recipient, rec):
        if not recipient:
            return
        key = (channel, recipient)
        cur = latest.get(key)
        if not cur or (rec.get("at") or "") >= (cur.get("at") or ""):
            latest[key] = rec

    async for r in db.iys_consents.find({}, {"_id": 0}).sort("created_at", 1):
        at = r.get("consent_date") or r.get("created_at") or ""
        st = "onay" if str(r.get("status") or "ONAY").upper() == "ONAY" else "ret"
        base = {"at": at, "status": st, "source": _SRC_TR.get(str(r.get("source") or ""), str(r.get("source") or "")),
                "origin": "İYS kaydı", "email": (r.get("email") or "").strip().lower(),
                "phone": r.get("phone") or "", "order_id": r.get("order_id") or "", "user_id": r.get("user_id"),
                "reported": bool(r.get("reported"))}
        for ch in (r.get("channels") or []):
            if ch == "EPOSTA" and base["email"]:
                _put("email", base["email"], dict(base, recipient=base["email"]))
            elif ch == "MESAJ" and base["phone"]:
                _put("sms", base["phone"], dict(base, recipient=base["phone"]))

    async for s in db.newsletter_subscribers.find({}, {"_id": 0}):
        em = (s.get("email") or "").strip().lower()
        if not em:
            continue
        at = s.get("unsubscribed_at") or s.get("consent_at") or s.get("created_at") or ""
        st = "onay" if (s.get("consent") and s.get("active", True) is not False) else "ret"
        _put("email", em, {"at": at, "status": st, "source": _SRC_TR.get(str(s.get("source") or "footer"), str(s.get("source") or "")),
                          "origin": "Bülten aboneliği", "email": em, "phone": "", "recipient": em,
                          "order_id": "", "user_id": None, "reported": True})

    # Üye profili pazarlama tercihi (Hesabım / admin üye kartı / sipariş izniyle True olur).
    # Kampanya kitlesiyle (email_marketing._campaign_audience) aynı kaynak — liste ile gönderim tutarlı.
    async for u in db.users.find({"accepts_marketing": True, "email": {"$exists": True, "$ne": ""}},
                                 {"_id": 0, "email": 1, "marketing_consent_at": 1, "created_at": 1, "id": 1}):
        em = (u.get("email") or "").strip().lower()
        if not em:
            continue
        at = u.get("marketing_consent_at") or u.get("created_at") or ""
        _put("email", em, {"at": at, "status": "onay", "source": _SRC_TR["account"], "origin": "Üye profili",
                          "email": em, "phone": "", "recipient": em, "order_id": "", "user_id": u.get("id"), "reported": True})

    suppressed = set()
    async for x in db.email_suppressions.find({}, {"_id": 0, "email": 1, "reason": 1}):
        em = (x.get("email") or "").strip().lower()
        if em:
            suppressed.add(em)
    for (ch, rcp), rec in latest.items():
        if ch == "email" and rcp in suppressed:
            rec["suppressed"] = True
    return latest


async def _names_for(recs: list) -> None:
    """Üye adını e-posta/telefon eşleşmesiyle doldur (tek sorgu)."""
    emails = {r["recipient"] for r in recs if r.get("channel") == "email"}
    phones = {r["recipient"] for r in recs if r.get("channel") == "sms"}
    by_email, by_phone = {}, {}
    if emails:
        async for u in db.users.find({"email": {"$in": list(emails)}}, {"_id": 0, "email": 1, "first_name": 1, "last_name": 1, "phone": 1}):
            by_email[(u.get("email") or "").lower()] = u
    if phones:
        # KVKK: üye telefonu şifreli → eşleşme HASH ile (regex/parça tarama yok); backfill öncesi düz
        # kayıtlar tam biçim eşleşmesiyle bulunur.
        from security.contact_pii import phone_hashes as _ph_hashes, legacy_phone_variants as _ph_vars, \
            normalize_phone as _ph_norm, open_field as _open
        tails = sorted({_ph_norm(p) for p in phones if len(_ph_norm(p)) == 10})
        if tails:
            h2t = {}
            variants = []
            for t in tails:
                for h in _ph_hashes(t):
                    h2t[h] = t
                variants += _ph_vars(t)
            async for u in db.users.find({"$or": [{"phone_hash": {"$in": list(h2t) or ["\x00"]}},
                                                   {"phone": {"$in": variants}}]},
                                         {"_id": 0, "email": 1, "first_name": 1, "last_name": 1, "phone": 1,
                                          "phone_enc": 1, "phone_hash": 1}):
                t = h2t.get(u.get("phone_hash") or "") or _ph_norm(_open(u, "phone"))
                if len(t) == 10:
                    by_phone[t] = u
    for r in recs:
        u = by_email.get(r["recipient"]) if r.get("channel") == "email" else by_phone.get(r["recipient"][-10:])
        if u:
            r["name"] = f"{u.get('first_name', '')} {u.get('last_name', '')}".strip()
            r["member_email"] = u.get("email")


def _rows(latest: dict, channel: Optional[str], status: Optional[str], q: str) -> list:
    out = []
    ql = (q or "").strip().lower()
    for (ch, rcp), rec in latest.items():
        if channel and ch != channel:
            continue
        st = "ret" if rec.get("suppressed") else rec.get("status")
        if status and st != status:
            continue
        if ql and ql not in rcp and ql not in (rec.get("email") or "") and ql not in (rec.get("phone") or ""):
            continue
        out.append({"channel": ch, "recipient": rcp, "status": st, "suppressed": bool(rec.get("suppressed")),
                    "at": rec.get("at"), "source": rec.get("source"), "origin": rec.get("origin"),
                    "email": rec.get("email"), "phone": rec.get("phone"), "reported": rec.get("reported")})
    out.sort(key=lambda r: r.get("at") or "", reverse=True)
    return out


@router.get("/summary")
async def consents_summary(current_user: dict = Depends(require_permission("customers.view", any_in_area=True))):
    latest = await _collect()
    def _n(ch, st):
        return sum(1 for (c, _), r in latest.items() if c == ch and (("ret" if r.get("suppressed") else r.get("status")) == st))
    return {"email": {"onay": _n("email", "onay"), "ret": _n("email", "ret")},
            "sms": {"onay": _n("sms", "onay"), "ret": _n("sms", "ret")},
            "generated_at": datetime.now(timezone.utc).isoformat()}


@router.get("")
async def list_consents(channel: Optional[str] = Query(None), status: Optional[str] = Query(None),
                        q: str = "", page: int = Query(1, ge=1), limit: int = Query(50, ge=1, le=500),
                        current_user: dict = Depends(require_permission("customers.view", any_in_area=True))):
    latest = await _collect()
    rows = _rows(latest, channel, status, q)
    total = len(rows)
    page_rows = rows[(page - 1) * limit: page * limit]
    await _names_for(page_rows)
    # KVKK: izin listesinde SMS alıcısı / telefon yalnız son 2 hane (tam liste: customers.export ile Excel).
    from security.contact_pii import mask_phone as _mask_phone
    for r in page_rows:
        if r.get("channel") == "sms":
            r["recipient"] = _mask_phone(r.get("recipient"))
        r["phone"] = _mask_phone(r.get("phone"))
    return {"items": page_rows, "total": total, "page": page, "pages": max(1, (total + limit - 1) // limit)}


# CSV/XLSX FORMÜL ENJEKSİYONU KORUMASI: kullanıcı kaynaklı metin (ad, kaynak, e-posta…) "=",
# "+", "-", "@", TAB, CR ile başlıyorsa Excel/LibreOffice onu formül olarak çalıştırabilir
# (=HYPERLINK, =cmd|…). Başına ' eklenir → düz metin. openpyxl "=" ile başlayan string'i
# doğrudan FORMÜL hücresi yapar; bu yüzden XLSX'te de zorunlu. Saf telefon/sayı ("+90 5xx…")
# formül yükü taşıyamaz → okunurluk için olduğu gibi bırakılır.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMERIC = re.compile(r"^[+\-]?[\d\s().\-]+$")


def _xls_safe(v):
    if not isinstance(v, str) or not v:
        return v
    if v.startswith(_FORMULA_PREFIXES) and not _PLAIN_NUMERIC.match(v):
        return "'" + v
    return v


@router.get("/export.xlsx")
async def export_consents(request: Request, channel: Optional[str] = Query(None), status: Optional[str] = Query(None),
                          q: str = "",
                          # GÜVENLİK: toplu kişisel veri (e-posta/telefon/ad) dışa aktarımı — üye dışa
                          # aktarımıyla AYNI yetki (customers.export); yalnız personel olmak yetmez.
                          current_user: dict = Depends(require_permission("customers.export"))):
    # ACİL ANAHTAR: freeze_exports açıkken toplu kişisel veri dışa aktarımı yapılmaz (423).
    from security_switches import require_switch_off, emit_security_event
    await require_switch_off("freeze_exports")
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    latest = await _collect()
    rows = _rows(latest, channel, status, q)
    await _names_for(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "İzinler"
    ws.append(["Kanal", "Alıcı", "Ad Soyad", "Durum", "Kaynak", "Kayıt türü", "Tarih", "E-posta", "Telefon", "İYS'ye bildirildi", "Kara liste"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="111111")
    for r in rows:
        ws.append([_xls_safe(c) for c in (
                   "E-posta" if r["channel"] == "email" else "SMS", r["recipient"], r.get("name") or "",
                   "İZİN VERDİ" if r["status"] == "onay" else "RET", r.get("source") or "", r.get("origin") or "",
                   str(r.get("at") or "")[:16].replace("T", " "), r.get("email") or "", r.get("phone") or "",
                   "Evet" if r.get("reported") else "Hayır", "Evet" if r.get("suppressed") else "")])
    for i, w in enumerate([10, 30, 24, 12, 18, 18, 17, 30, 16, 12, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"pazarlama-izinleri-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.xlsx"
    # KVKK: kişisel veri içeren toplu dışa aktarım denetim kaydına düşer (kim/ne zaman/hangi
    # filtre/kaç satır). Arama metni PII içerebileceği için yalnız "var/yok" yazılır.
    await record_admin_audit(
        db, action="consents.export", entity_type="consents", entity_id="export.xlsx",
        current_user=current_user, request=request,
        metadata={"channel": channel or "all", "status": status or "all",
                  "search_used": bool((q or "").strip()), "row_count": len(rows)},
    )
    await emit_security_event("consent_export", current_user, {"export": "consents.xlsx", "rows": len(rows)})
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}", "Cache-Control": "no-store"})
